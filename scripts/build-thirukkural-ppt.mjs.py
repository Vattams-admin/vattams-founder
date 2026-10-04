from pathlib import Path
import json, urllib.request, textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DATA_URL = "https://raw.githubusercontent.com/tk120404/thirukkural/master/thirukkural.json"
DETAIL_URL = "https://raw.githubusercontent.com/tk120404/thirukkural/master/detail.json"
OUT = Path("artifacts/VATTAMS_Academia_Thirukkural_1330_Full.pptx")

NAVY = RGBColor(7, 27, 55)
NAVY2 = RGBColor(12, 45, 82)
GOLD = RGBColor(205, 154, 55)
CREAM = RGBColor(250, 244, 228)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 35, 42)
GREEN = RGBColor(38, 112, 70)
PURPLE = RGBColor(86, 57, 125)

def fetch(url):
    with urllib.request.urlopen(url, timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))

def add_bg(slide, title=None, chapter_no=None, total=None):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    # gold top/bottom accents
    for y, h in [(0, 0.08), (7.38, 0.12)]:
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(y), Inches(13.333), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = GOLD
        sh.line.fill.background()
    if title:
        tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(9.8), Inches(0.55))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.name = "Noto Sans Tamil"; p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = WHITE
    if chapter_no is not None:
        tb = slide.shapes.add_textbox(Inches(10.25), Inches(0.2), Inches(2.5), Inches(0.45))
        p = tb.text_frame.paragraphs[0]
        p.text = f"அதிகாரம் {chapter_no}  •  {total}"
        p.alignment = PP_ALIGN.RIGHT
        p.font.name = "Noto Sans Tamil"; p.font.size = Pt(12); p.font.color.rgb = GOLD

def add_text(slide, x, y, w, h, text, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Pt(7); tf.margin_right = Pt(7); tf.margin_top = Pt(4); tf.margin_bottom = Pt(4)
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = "Noto Sans Tamil"; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return tb

def add_card(slide, k, x, y, w, h):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = CREAM
    card.line.color.rgb = GOLD; card.line.width = Pt(1.2)
    # number
    add_text(slide, x+0.08, y+0.05, 0.7, 0.3, f"குறள் {k['Number']}", 11, GOLD, True)
    # verse
    verse = f"{k['Line1']}\n{k['Line2']}"
    add_text(slide, x+0.08, y+0.36, w-0.16, 0.78, verse, 16, NAVY, True)
    # meaning and explanation
    mv = k.get("mv") or ""
    sp = k.get("sp") or ""
    add_text(slide, x+0.08, y+1.15, w-0.16, 0.92, "பொருள்: " + mv, 11.5, INK, False)
    add_text(slide, x+0.08, y+2.03, w-0.16, h-2.12, "விளக்கம்: " + sp, 10.8, PURPLE, False)

data = fetch(DATA_URL)["kural"]
detail = fetch(DETAIL_URL)
chapters = []
def walk(node):
    if isinstance(node, dict):
        if "chapters" in node and isinstance(node["chapters"], dict):
            for ch in node["chapters"].get("detail", []):
                chapters.append(ch)
        for v in node.values():
            if isinstance(v, (dict, list)): walk(v)
    elif isinstance(node, list):
        for v in node: walk(v)
walk(detail)
# dedupe and sort by chapter number
chapters = {c["number"]: c for c in chapters}
chapters = [chapters[i] for i in sorted(chapters)]
by_no = {k["Number"]: k for k in data}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Cover
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
add_text(s, 0.8, 0.75, 11.7, 0.8, "VATTAMS", 40, WHITE, True, PP_ALIGN.CENTER)
add_text(s, 0.8, 1.5, 11.7, 0.7, "ACADEMIA", 30, GOLD, True, PP_ALIGN.CENTER)
add_text(s, 0.8, 2.35, 11.7, 0.9, "திருக்குறள்", 42, GOLD, True, PP_ALIGN.CENTER)
add_text(s, 0.8, 3.15, 11.7, 0.55, "133 அதிகாரங்கள் • 1330 குறள்கள்", 21, WHITE, True, PP_ALIGN.CENTER)
add_text(s, 1.3, 4.0, 10.7, 1.0, "குறள் • பொருள் • விளக்கம்", 28, CREAM, True, PP_ALIGN.CENTER)
add_text(s, 1.3, 5.2, 10.7, 0.55, "அறத்துப்பால் • பொருட்பால் • காமத்துப்பால்", 17, WHITE, False, PP_ALIGN.CENTER)
add_text(s, 1.3, 6.25, 10.7, 0.4, "EMPOWERING MINDS. BUILDING FUTURES.", 12, GOLD, True, PP_ALIGN.CENTER)

# Section divider helper
def section_slide(title, subtitle, color=GOLD):
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
    add_text(s, 0.7, 2.1, 11.9, 0.8, title, 34, color, True, PP_ALIGN.CENTER)
    add_text(s, 1.0, 3.1, 11.3, 0.7, subtitle, 20, WHITE, False, PP_ALIGN.CENTER)
    return s

# Create 2 content slides per chapter, 5 kurals per slide
for ch in chapters:
    start, end = ch["start"], ch["end"]
    ks = [by_no[n] for n in range(start, end+1) if n in by_no]
    for page_idx in range(2):
        chunk = ks[page_idx*5:(page_idx+1)*5]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_bg(s, f"{ch['number']}. {ch['name']}", ch["number"], f"{start}–{end}")
        add_text(s, 0.58, 0.78, 12.15, 0.35, f"{ch.get('translation','')}  •  {ch.get('transliteration','')}", 10.5, GOLD, False, PP_ALIGN.CENTER)
        y = 1.18
        for k in chunk:
            add_card(s, k, 0.58, y, 12.15, 1.18)
            y += 1.23
        add_text(s, 0.58, 7.08, 12.15, 0.2, "VATTAMS ACADEMIA  •  திருக்குறள் கற்றல் தொகுப்பு", 8.5, WHITE, False, PP_ALIGN.CENTER)

# Sources / notes
s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
add_text(s, 0.7, 0.65, 11.9, 0.6, "ஆதாரங்கள் மற்றும் பயன்பாட்டு குறிப்பு", 28, GOLD, True, PP_ALIGN.CENTER)
notes = (
    "• குறள் மூல உரை, அதிகார அமைப்பு மற்றும் குறள் எண்கள்: tk120404/thirukkural public dataset.\n"
    "• 'பொருள்' பகுதியில் உள்ள உரை: dataset-இல் உள்ள 'mv' — டாக்டர் மு. வரதராசனார் உரை.\n"
    "• 'விளக்கம்' பகுதியில் உள்ள உரை: dataset-இல் உள்ள 'sp' — சாலமன் பாப்பையா உரை.\n"
    "• திருக்குறள்: 3 பால், 133 அதிகாரங்கள், 1330 குறள்கள்.\n"
    "• இந்த PPT கல்வி / கற்றல் பயன்பாட்டிற்காக தொகுக்கப்பட்டுள்ளது. உரை ஆதாரங்களுக்கு உரிய ஆசிரியர்/தரவு ஆதார மரியாதை வழங்கவும்."
)
add_text(s, 1.0, 1.65, 11.3, 3.7, notes, 17, CREAM, False)
add_text(s, 1.0, 5.75, 11.3, 0.75, "Source dataset: https://github.com/tk120404/thirukkural", 12, GOLD, False, PP_ALIGN.CENTER)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Created {OUT} with {len(prs.slides)} slides.")
